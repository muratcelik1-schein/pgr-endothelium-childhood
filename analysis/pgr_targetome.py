#!/usr/bin/env python3
"""Does the receptor's binding repertoire predict co-regulation in human brain endothelium?

The binding set comes from GSE43786, chromatin immunoprecipitation of the
progesterone receptor in primary human umbilical vein endothelial cells carrying
the receptor from a lentivirus and treated with progesterone. Two things about
that source travel with every number below and are written into the output:
the cells are venous endothelium but not brain, and the receptor is introduced
rather than endogenous.

Only the binding-based set is used for the main test, because peaks were called
against three negative controls and need no replication. The accompanying
RNA-seq of that study has one sample per condition, so the up and down subsets
derived from it are tested separately and labelled as resting on unreplicated
expression data.

The test. In the postnatal endothelial series, each gene's partial Spearman
correlation with PGR is computed at donor level with age removed, so that a
gene rising with age for its own reasons does not score. The bound set is then
compared with sets of non-bound genes matched on expression abundance, drawn
many times. Matching on abundance matters because detection rate drives
correlation in single nucleus data.

Controls. A permutation that shuffles the PGR vector bounds the false positive
rate. The abundance-matched draw bounds the effect of detection. Both are
reported whether or not they behave.

Run: python3 pgr_targetome.py --h5ad <consensus atlas> --peaks <peaks.xlsx> --out out/targetome.tsv
"""
import argparse, os, re, sys
import numpy as np
import h5py
from scipy import stats

SEED = 2026
rng = np.random.default_rng(SEED)
TARGET = "PGR"
AGE_MIN, MIN_CELL, MIN_DET = 0.75, 10, 0.20
N_DRAW = 2000
N_BIN = 20

ap = argparse.ArgumentParser()
ap.add_argument("--h5ad", required=True)
ap.add_argument("--peaks", required=True, help="GSM1071297 ...peaks.xlsx")
ap.add_argument("--out", required=True)
A = ap.parse_args()

rows = []
def add(b, k, v, n=""):
    rows.append((b, k, v, n)); print(f"{b:20s} {k:36s} {v}   {n}", flush=True)
def flush():
    os.makedirs(os.path.dirname(A.out), exist_ok=True)
    with open(A.out, "w") as fh:
        fh.write("block\tkey\tvalue\tnote\n")
        for r in rows: fh.write("\t".join(str(x) for x in r) + "\n")

# ------------------------------------------------------------------ gene sets
import openpyxl
wb = openpyxl.load_workbook(A.peaks, read_only=True)
add("source", "workbook", os.path.basename(A.peaks))
add("source", "sheets", "; ".join(wb.sheetnames))
def sheet_genes(name_part, col=0):
    hit = [s for s in wb.sheetnames if name_part.lower() in s.lower()]
    assert len(hit) == 1, f"one sheet was expected for {name_part!r}: {hit}"
    ws = wb[hit[0]]
    g = []
    for r in ws.iter_rows(values_only=True):
        if not r or r[col] is None: continue
        v = str(r[col]).strip()
        if not v or v.lower() in ("gene",) or v.startswith("#") or " " in v: continue
        g.append(v)
    return hit[0], sorted(set(g))
sn_b, BOUND = sheet_genes("Bound genes nearest")
sn_d, DOWN = sheet_genes("Bound_downreg")
sn_u, UP = sheet_genes("Bound_upreg")
for nm, s, g in [("bound", sn_b, BOUND), ("bound_down", sn_d, DOWN), ("bound_up", sn_u, UP)]:
    add("gene_sets", nm, len(g), f"sheet {s!r}")
npk = len([r for r in wb["FinalList_C2vsC3andC4andC6-neg"].iter_rows(values_only=True) if r and r[0]])
add("gene_sets", "peaks_passing_all_negative_controls", npk, "hg19, MACS, GREAT 50 kb nearest gene")

# ------------------------------------------------------------------ endothelium
f = h5py.File(A.h5ad, "r"); obs = f["obs"]
def col(n):
    g = obs[n]
    if isinstance(g, h5py.Group) and "categories" in g:
        c = np.array([x.decode() if isinstance(x, bytes) else str(x) for x in g["categories"][:]], dtype=object)
        return c[g["codes"][:]]
    v = g[:]
    return np.array([x.decode() if isinstance(x, bytes) else str(x) for x in v], dtype=object)
cc, fa, dn, st = col("cellclass"), col("final_annotation"), col("donor_id"), col("study")
age = obs["age_years_postgestation"][:].astype(float)
keep = (cc == "Vascular") & (fa == "Endo") & (age > AGE_MIN)
sym = np.array([x.decode() if isinstance(x, bytes) else str(x) for x in f["var"]["gene_symbol"][:]], dtype=object)
G = len(sym)
X = f["X"]; d0 = X["data"][:200000]
assert np.all(d0 == np.round(d0)), "integer counts are required"
indptr = X["indptr"][:]; data, indices = X["data"], X["indices"]

sel = np.where(keep)[0]
donors = sorted(set(dn[sel]))
di = {d: i for i, d in enumerate(donors)}
PB = np.zeros((len(donors), G), dtype=np.float64)
NC = np.zeros(len(donors), dtype=int)
for r in sel:
    a, b = indptr[r], indptr[r + 1]
    if b > a: np.add.at(PB[di[dn[r]]], indices[a:b], data[a:b])
    NC[di[dn[r]]] += 1
ok = NC >= MIN_CELL
PB, NC = PB[ok], NC[ok]
dkeep = [d for d, k in zip(donors, ok) if k]
dage = np.array([age[sel][list(dn[sel]).index(d)] for d in dkeep])
dstudy = np.array([st[sel][list(dn[sel]).index(d)] for d in dkeep])
add("series", "donors", len(dkeep), "manuscript reports 65")
add("series", "nuclei", int(NC.sum()), "manuscript reports 5,388")

CP = PB / np.maximum(PB.sum(1, keepdims=True), 1) * 1e4
det = (CP > 0).mean(0)
gkeep = det >= MIN_DET
CP, sym_k = CP[:, gkeep], sym[gkeep]
add("series", "genes_tested", int(gkeep.sum()), "manuscript reports 18,794")
Z = np.log1p(CP)
for s in set(dstudy):
    j = dstudy == s
    Z[j] = (Z[j] - Z[j].mean(0)) / np.maximum(Z[j].std(0), 1e-9)

pos = {g: i for i, g in enumerate(sym_k)}
assert TARGET in pos, "PGR absent from the series"

# ------------------------------------------------------------------ partial correlation with PGR, age removed
def resid(y, C):
    C = np.column_stack([np.ones(len(y))] + [c for c in C])
    return y - C @ np.linalg.lstsq(C, y, rcond=None)[0]
R = lambda v: stats.rankdata(v)
a_r = R(dage)
pgr_r = resid(R(Z[:, pos[TARGET]]), [a_r])
Zr = np.apply_along_axis(R, 0, Z)
Zres = Zr - np.column_stack([np.ones(len(dage)), a_r]) @ np.linalg.lstsq(
    np.column_stack([np.ones(len(dage)), a_r]), Zr, rcond=None)[0]
num = pgr_r @ Zres
den = np.sqrt((pgr_r @ pgr_r) * (Zres * Zres).sum(0))
rho = num / np.maximum(den, 1e-12)
add("method", "statistic", "partial Spearman with PGR, age removed, donor level")

# ------------------------------------------------------------------ abundance-matched null
mean_ab = np.log1p(CP).mean(0)
bins = np.clip(np.digitize(mean_ab, np.quantile(mean_ab, np.linspace(0, 1, N_BIN + 1)[1:-1])), 0, N_BIN - 1)
def matched_draw(idx):
    out = []
    for b in range(N_BIN):
        need = int((bins[idx] == b).sum())
        if not need: continue
        pool = np.setdiff1d(np.where(bins == b)[0], idx)
        if len(pool) == 0: continue
        out.append(rng.choice(pool, size=min(need, len(pool)), replace=False))
    return np.concatenate(out) if out else np.array([], dtype=int)

def test(name, genes, note=""):
    idx = np.array([pos[g] for g in genes if g in pos and g != TARGET])
    if len(idx) < 20:
        add("result", name, f"only {len(idx)} genes present in the series", "not tested"); return
    obs_stat = float(np.median(np.abs(rho[idx])))
    null = np.array([float(np.median(np.abs(rho[matched_draw(idx)]))) for _ in range(N_DRAW)])
    p = float(((null >= obs_stat).sum() + 1) / (N_DRAW + 1))
    obs_signed = float(np.median(rho[idx]))
    null_s = np.array([float(np.median(rho[matched_draw(idx)])) for _ in range(200)])
    add("result", name,
        f"genes={len(idx)}/{len(genes)}; median |partial rho|={obs_stat:.4f}; "
        f"null median={np.median(null):.4f}; p={p:.4f}; signed median={obs_signed:+.4f} "
        f"(null {np.median(null_s):+.4f})",
        note or "abundance matched null, 2000 draws")
    return p

add("controls", "positive_control_genes",
    f"PGR-AS1 rho={rho[pos['PGR-AS1']]:+.3f}; ZNF385B rho={rho[pos['ZNF385B']]:+.3f}"
    if "PGR-AS1" in pos and "ZNF385B" in pos else "absent",
    "genes already reported as co-rising; the statistic should see them")

test("bound_set", BOUND)
test("bound_and_downregulated", DOWN, "rests on unreplicated RNA-seq, one sample per condition")
test("bound_and_upregulated", UP, "rests on unreplicated RNA-seq, one sample per condition")

# shuffled-target control
sh = []
for _ in range(200):
    perm = rng.permutation(len(dage))
    pr = resid(R(Z[perm, pos[TARGET]]), [a_r])
    n2 = pr @ Zres; d2 = np.sqrt((pr @ pr) * (Zres * Zres).sum(0))
    r2 = n2 / np.maximum(d2, 1e-12)
    idx = np.array([pos[g] for g in BOUND if g in pos and g != TARGET])
    sh.append(float(np.median(np.abs(r2[idx]))))
idx = np.array([pos[g] for g in BOUND if g in pos and g != TARGET])
add("controls", "shuffled_target_null",
    f"observed {np.median(np.abs(rho[idx])):.4f}; shuffled median {np.median(sh):.4f}; "
    f"p={((np.array(sh) >= np.median(np.abs(rho[idx]))).sum()+1)/201:.4f}",
    "PGR vector permuted across donors, 200 times")

add("limits", "binding_source", "primary human umbilical vein endothelium, receptor introduced by lentivirus",
    "venous endothelium but not brain, and not endogenous receptor")
add("limits", "rna_source", "one sample per condition in the accompanying RNA-seq",
    "the up and down subsets are candidate lists, not differential expression results")
add("limits", "assembly", "peaks are hg19, genes assigned by GREAT 2.0.1, nearest gene within 50 kb")
add("limits", "seed", SEED)
flush()
print(f"\nwritten: {A.out} ({len(rows)} rows)")
