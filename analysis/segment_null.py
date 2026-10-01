#!/usr/bin/env python3
"""Two things that were set aside, tested properly instead of assumed.

ONE. The effector modules correlate with the receptor across donors once age is
removed, at +0.41 for the binding set against a matched null of +0.28. That was
set aside on the grounds that a donor quality axis could produce it. But the
housekeeping module did not show the same correlation, which argues the other
way, so the question was never actually settled. It is settled here by the
control that was missing: the same modules are also correlated with endothelial
HOUSEKEEPING expression, with age removed. A module that tracks the receptor and
not the housekeeping panel is telling us something; a module that tracks both is
telling us about the donors.

TWO. The venous preference of the transcript rests on a three class assignment
and on a chi square over nuclei. A reanalysis of a different cohort found that
against abundance matched genes the venular preference sits at p = 0.055, with
more than half of matched genes moving the same way. That test has never been
run in the primary series. It is run here: how many abundance matched genes show
a venous preference as large as the target's, in the same nuclei, with the same
assignment.

Neither test is new data. Both are the control that the earlier round skipped.

Run: python3 recheck.py --h5ad F --peaks P.xlsx --perm-map M.tsv --out T
"""
import argparse, os
import numpy as np
import h5py
from scipy import stats

SEED = 2026
rng = np.random.default_rng(SEED)
TARGET = "PGR"
AGE_MIN, MIN_CELL, MIN_DET = 0.75, 10, 0.20
N_DRAW, N_BIN = 2000, 20
SEGMENT = {"arterial":  ["GJA5", "SEMA3G", "HEY1", "ALPL", "VEGFC", "BMX"],
           "capillary": ["MFSD2A", "SLC7A5", "TFRC", "SLC16A1", "RGCC"],
           "venous":    ["NR2F2", "VWF", "ACKR1", "PLVAP", "IL1R1"]}
HOUSE = ["ACTB", "GAPDH", "RPL13A", "RPLP0", "TBP", "PGK1", "B2M", "PPIA", "UBC",
         "YWHAZ", "SDHA", "HPRT1", "RPL19", "RPS18", "EEF1A1", "PSMB4", "VPS29",
         "CHMP2A", "EMC7", "GPI", "REEP5", "SNRPD3", "VCP", "RAB7A"]

ap = argparse.ArgumentParser()
for a in ("--h5ad", "--peaks", "--perm-map", "--out"): ap.add_argument(a, required=True)
A = ap.parse_args()
rows = []
def add(b, k, v, n=""):
    rows.append((b, k, str(v), n)); print(f"{b:16s} {k:34s} {v}   {n}", flush=True)

import openpyxl
wb = openpyxl.load_workbook(A.peaks, read_only=True)
def sheet(part):
    hit = [s for s in wb.sheetnames if part.lower() in s.lower()]
    assert len(hit) == 1, hit
    out = []
    for r in wb[hit[0]].iter_rows(values_only=True):
        if not r or r[0] is None: continue
        v = str(r[0]).strip()
        if v and v.lower() != "gene" and " " not in v and not v.startswith("#"): out.append(v)
    return sorted(set(out))
BOUND = sheet("Bound genes nearest")
PERM = sorted({ln.split("\t")[1].strip() for ln in open(A.perm_map).read().splitlines()[1:] if "\t" in ln})

f = h5py.File(A.h5ad, "r"); obs = f["obs"]
def col(n):
    g = obs[n]
    if isinstance(g, h5py.Group) and "categories" in g:
        c = np.array([x.decode() if isinstance(x, bytes) else str(x) for x in g["categories"][:]], dtype=object)
        return c[g["codes"][:]]
    return np.array([x.decode() if isinstance(x, bytes) else str(x) for x in g[:]], dtype=object)
cc, fa, dn, st = col("cellclass"), col("final_annotation"), col("donor_id"), col("study")
age = obs["age_years_postgestation"][:].astype(float)
sym = np.array([x.decode() if isinstance(x, bytes) else str(x) for x in f["var"]["gene_symbol"][:]], dtype=object)
X = f["X"]; indptr = X["indptr"][:]; data, indices = X["data"], X["indices"]
sel = np.where((cc == "Vascular") & (fa == "Endo") & (age > AGE_MIN))[0]
ds = sorted(set(dn[sel])); di = {d: i for i, d in enumerate(ds)}
PB = np.zeros((len(ds), len(sym))); NC = np.zeros(len(ds), int); CELLS = {d: [] for d in ds}; FIRST = {}
for r in sel:
    a_, b_ = indptr[r], indptr[r + 1]
    if b_ > a_: np.add.at(PB[di[dn[r]]], indices[a_:b_], data[a_:b_])
    NC[di[dn[r]]] += 1; CELLS[dn[r]].append(r); FIRST.setdefault(dn[r], r)
ok = NC >= MIN_CELL
PB, NC = PB[ok], NC[ok]; dk = [d for d, k in zip(ds, ok) if k]
dage = np.array([age[FIRST[d]] for d in dk]); dstudy = np.array([st[FIRST[d]] for d in dk])
add("gate", "series", f"{len(dk)} donors, {int(NC.sum())} nuclei", "manuscript reports 65 and 5,388")
assert len(dk) == 65 and int(NC.sum()) == 5388
CP = PB / np.maximum(PB.sum(1, keepdims=True), 1) * 1e4
gk = (CP > 0).mean(0) >= MIN_DET
CPk, symk = CP[:, gk], sym[gk]
Z = np.log1p(CPk)
for u in set(dstudy):
    j = dstudy == u
    Z[j] = (Z[j] - Z[j].mean(0)) / np.maximum(Z[j].std(0), 1e-9)
pos = {g: i for i, g in enumerate(symk)}
pgr = Z[:, pos[TARGET]]
house = Z[:, [pos[g] for g in HOUSE if g in pos]].mean(1)
ab = np.log1p(CPk).mean(0)
bins = np.clip(np.digitize(ab, np.quantile(ab, np.linspace(0, 1, N_BIN + 1)[1:-1])), 0, N_BIN - 1)
def matched(idx):
    out = []
    for b in range(N_BIN):
        need = int((bins[idx] == b).sum())
        if not need: continue
        pool = np.setdiff1d(np.where(bins == b)[0], idx)
        if len(pool): out.append(rng.choice(pool, size=min(need, len(pool)), replace=False))
    return np.concatenate(out) if out else np.array([], int)
def partial(x, y, c):
    R = stats.rankdata
    M = np.column_stack([np.ones(len(c)), R(c)])
    rx = R(x) - M @ np.linalg.lstsq(M, R(x), rcond=None)[0]
    ry = R(y) - M @ np.linalg.lstsq(M, R(y), rcond=None)[0]
    return float(rx @ ry / max(np.sqrt((rx @ rx) * (ry @ ry)), 1e-12))

# ---------------------------------------------------------------- ONE
add("test_one", "question", "does the module track the receptor, or the donor?")
for nm, genes in [("bound", BOUND), ("permeability", PERM), ("housekeeping_control", HOUSE)]:
    idx = np.array([pos[g] for g in genes if g in pos and g != TARGET], dtype=int)
    if len(idx) < 10: continue
    s = Z[:, idx].mean(1)
    r_pgr, r_house = partial(s, pgr, dage), partial(s, house, dage)
    np_, nh = [], []
    for _ in range(N_DRAW):
        j = matched(idx)
        if len(j) < 10: continue
        sn = Z[:, j].mean(1)
        np_.append(partial(sn, pgr, dage)); nh.append(partial(sn, house, dage))
    np_, nh = np.array(np_), np.array(nh)
    add("test_one", nm,
        f"with_PGR={r_pgr:+.3f} (null {np.median(np_):+.3f}, p={((np.abs(np_)>=abs(r_pgr)).sum()+1)/(len(np_)+1):.4f}); "
        f"with_endothelial_housekeeping={r_house:+.3f} (null {np.median(nh):+.3f}, p={((np.abs(nh)>=abs(r_house)).sum()+1)/(len(nh)+1):.4f})",
        "a module that tracks both is tracking the donor")

# ---------------------------------------------------------------- TWO
symidx = {g: i for i, g in enumerate(sym)}
nuc, nucdon = [], []
for d in dk: nuc += CELLS[d]; nucdon += [d] * len(CELLS[d])
nuc = np.array(nuc)
want = sorted({symidx[g] for v in SEGMENT.values() for g in v if g in symidx} |
              {symidx[g] for g in symk if g in symidx})
wpos = {g: i for i, g in enumerate(want)}; wset = set(want)
M = np.zeros((len(nuc), len(want)), dtype=np.float32); tot = np.zeros(len(nuc))
for n, r in enumerate(nuc):
    a_, b_ = indptr[r], indptr[r + 1]
    if b_ <= a_: continue
    ii, dd = indices[a_:b_], data[a_:b_]
    tot[n] = dd.sum()
    for j, val in zip(ii, dd):
        if int(j) in wset: M[n, wpos[int(j)]] = val
CPn = M / np.maximum(tot, 1)[:, None] * 1e4
def zc(cols):
    v = np.log1p(CPn[:, [wpos[symidx[g]] for g in cols]])
    v = (v - v.mean(0)) / np.maximum(v.std(0), 1e-9)
    return v.mean(1)
S = np.column_stack([zc([g for g in SEGMENT[k] if g in symidx]) for k in ("arterial", "capillary", "venous")])
assign = np.array(["arterial", "capillary", "venous"])[S.argmax(1)]
for k in ("arterial", "capillary", "venous"):
    add("segment", k, int((assign == k).sum()))
def venous_pref(gidx):
    """share of a gene's positive nuclei that are venous, minus the same for arterial"""
    v = M[:, wpos[gidx]] > 0
    if v.sum() < 20: return None
    return float((assign[v] == "venous").mean() - (assign[v] == "arterial").mean())
tgt = venous_pref(symidx[TARGET])
add("test_two", "target_venous_minus_arterial", f"{tgt:+.4f}",
    "among nuclei positive for the gene, venous share minus arterial share")
cand = [g for g in symk if g in symidx and g != TARGET]
tab = np.array([bins[pos[TARGET]]])
pool_idx = [pos[g] for g in cand]
same_bin = [g for g in cand if bins[pos[g]] == bins[pos[TARGET]]]
vals = []
for g in same_bin:
    v = venous_pref(symidx[g])
    if v is not None: vals.append(v)
vals = np.array(vals)
if len(vals) >= 30:
    pct = float((vals < tgt).mean() * 100)
    p = float(((vals >= tgt).sum() + 1) / (len(vals) + 1))
    add("test_two", "against_abundance_matched_genes",
        f"{len(vals)} genes in the target's abundance bin; target at the {pct:.1f}th percentile; p={p:.4f}",
        "the control the earlier segment analysis never ran")
    add("test_two", "matched_gene_median", f"{np.median(vals):+.4f}")
else:
    add("test_two", "against_abundance_matched_genes", f"only {len(vals)} matched genes", "not tested")
add("limit", "seed", SEED)
os.makedirs(os.path.dirname(A.out), exist_ok=True)
with open(A.out, "w") as fh:
    fh.write("block\tkey\tvalue\tnote\n")
    for r in rows: fh.write("\t".join(r) + "\n")
print(f"\nwritten: {A.out} ({len(rows)} rows)")
